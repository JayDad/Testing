"""
motion_slides.py — 장면(Scene) YAML 1개 → ① 모핑 PPTX  ② 애니메이션 HTML  ③ MP4

설계 원칙
- 모든 개체는 'box' 1종(위치·크기·채움·모서리·그림자·텍스트 줄)으로 단순화
- 장면 간 같은 id = 같은 개체 → PPTX는 '!!id' 이름으로 모핑 강제 매칭, HTML은 보간
- 장면은 직전 장면을 상속(변경분만 기술), remove/reset으로 제거·초기화
- 좌표 단위 inch (16:9, 13.333 × 7.5), 글자 크기 pt — 다크블루 테마 좌표계와 동일
"""

from __future__ import annotations

import base64
import copy
import json
import os
import shutil
import subprocess
from pathlib import Path

import yaml
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
FONT_FILE = ROOT / "fonts" / "PretendardVariable.woff2"

SLIDE_W, SLIDE_H = 13.333, 7.5
PX_PER_IN = 144  # 1920 × 1080 캔버스
FONT = "Pretendard"

# ---------------------------------------------------------------------------
# Design tokens (다크블루 테마)
# ---------------------------------------------------------------------------
TOKENS = {
    "navy": "012060",
    "text": "404040",
    "bg": "F2F2F2",
    "white": "FFFFFF",
    "highlight": "E6EBF5",
    "divider": "BFBFBF",
    "caption": "808080",
}

# 스타일 프리셋 — 개체에 style: 로 지정, 개별 키로 덮어쓰기 가능
STYLES = {
    "text":      dict(fill=None, radius=0, shadow=False, pad=0.05, color="text"),
    "card":      dict(fill="white", radius=0.10, shadow=True, pad=0.25, color="text"),
    "navycard":  dict(fill="navy", radius=0.10, shadow=True, pad=0.25, color="white"),
    "highlight": dict(fill="highlight", radius=0.10, shadow=False, pad=0.25, color="text"),
    "banner":    dict(fill="navy", radius=0.06, shadow=True, pad=0.15, color="white",
                      align="center", anchor="middle"),
    "bar":       dict(fill="navy", radius=0, shadow=False, pad=0, color="white"),
}

BOX_DEFAULTS = dict(x=0.6, y=1.85, w=3.0, h=1.0, fill=None, radius=0, shadow=False,
                    pad=0.05, align="left", anchor="top", opacity=1.0, color="text",
                    size=14, bold=False, lines=[])


def color_hex(c):
    if c is None or c == "none":
        return None
    c = str(c)
    return TOKENS.get(c, c.lstrip("#")).upper()


# ---------------------------------------------------------------------------
# Spec 해석 — 상속·프리셋·크롬(제목/구분선/페이지번호) 전개
# ---------------------------------------------------------------------------

def _norm_lines(raw, box):
    """text: '문자열' | lines: [문자열 | {t,size,bold,color,gap}] → 표준 줄 목록"""
    out = []
    for ln in raw:
        if isinstance(ln, str):
            ln = {"t": ln}
        out.append(dict(
            t=str(ln.get("t", "")),
            size=float(ln.get("size", box["size"])),
            bold=bool(ln.get("bold", box["bold"])),
            color=color_hex(ln.get("color", box["color"])),
            gap=float(ln.get("gap", 4)),  # 단락 뒤 간격(pt)
        ))
    return out


def _resolve_box(eid, raw):
    box = dict(BOX_DEFAULTS)
    box.update(STYLES.get(raw.get("style", "text"), {}))
    box.update({k: v for k, v in raw.items() if k not in ("text", "lines", "style")})
    if "text" in raw:
        lines = [raw["text"]] if isinstance(raw["text"], str) else raw["text"]
    else:
        lines = raw.get("lines", [])
    box["lines"] = _norm_lines(lines, box)
    box["fill"] = color_hex(box["fill"])
    box["id"] = eid
    return box


def _chrome(meta, scene, idx, total):
    """다크블루 고정 크롬 — 장면 간 동일 id라 모핑 시 정지 상태 유지"""
    els = {}
    if meta.get("project_title"):
        els["_proj"] = dict(x=10.13, y=0.50, w=2.60, h=0.30, size=11, align="right",
                            text=meta["project_title"])
    if scene.get("title"):
        els["_title"] = dict(x=0.60, y=0.45, w=9.40, h=0.70, size=34, bold=True,
                             color="navy", anchor="bottom", text=scene["title"])
    if scene.get("subtitle"):
        sub = scene["subtitle"]
        sub = sub if sub.lstrip().startswith("—") else "— " + sub
        els["_sub"] = dict(x=0.60, y=1.15, w=9.40, h=0.42, size=20, text=sub)
    els["_div"] = dict(style="bar", x=0.60, y=1.645, w=12.13, h=0.014)
    if meta.get("page_numbers", True):
        els["_pg"] = dict(x=11.23, y=7.00, w=1.50, h=0.30, size=9, bold=True,
                          color="navy", align="right", text=f"{idx:02d}  /  {total:02d}")
    return els


def load_spec(path):
    spec = yaml.safe_load(Path(path).read_text(encoding="utf-8"))
    meta = dict(fps=30, hold=2.5, morph=1.2, intro=0.8, outro=0.0, chrome=True,
                page_numbers=True, project_title="", logo=None)
    meta.update(spec.get("meta") or {})
    raw_scenes = spec["scenes"]
    total = len(raw_scenes)

    scenes, current, last = [], {}, {}
    for i, sc in enumerate(raw_scenes, 1):
        sc = dict(sc)
        for k in ("title", "subtitle"):  # 제목·부제 미기재 시 직전 장면 유지
            sc.setdefault(k, last.get(k, ""))
            last[k] = sc[k]
        if sc.get("reset"):
            current = {}
        for rid in sc.get("remove", []) or []:
            current.pop(rid, None)
        for el in sc.get("elements", []) or []:
            el = dict(el)
            eid = str(el.pop("id"))
            if el.pop("replace", False) or eid not in current:
                current[eid] = el
            else:  # 상속: 바뀐 키만 덮어씀
                merged = dict(current[eid])
                merged.update(el)
                if "text" in el:
                    merged.pop("lines", None)
                if "lines" in el:
                    merged.pop("text", None)
                current[eid] = merged
        raw_all = dict(_chrome(meta, sc, i, total)) if meta["chrome"] else {}
        raw_all.update(current)
        boxes = [_resolve_box(eid, copy.deepcopy(r)) for eid, r in raw_all.items()]
        scenes.append(dict(
            title=sc.get("title", ""),
            hold=float(sc.get("hold", meta["hold"])),
            morph=float(sc.get("morph", meta["morph"] if i > 1 else meta["intro"])),
            notes=sc.get("notes", ""),
            boxes=boxes,
        ))
    return meta, scenes


# ---------------------------------------------------------------------------
# ① PPTX + 모핑 전환
# ---------------------------------------------------------------------------
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS_P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def _transition_xml(kind, dur_ms, adv_ms):
    """모핑 전환(미지원 버전은 페이드로 대체 재생) — python-pptx 미지원 영역이라 XML 직접 삽입"""
    effect = '<p159:morph option="byObject"/>' if kind == "morph" else "<p:fade/>"
    return etree.fromstring(
        f'<mc:AlternateContent xmlns:mc="{NS_MC}" xmlns:p="{NS_P}">'
        f'<mc:Choice xmlns:p14="{NS_P14}" xmlns:p159="{NS_P159}" Requires="p159">'
        f'<p:transition spd="slow" p14:dur="{dur_ms}" advTm="{adv_ms}">{effect}</p:transition>'
        f'</mc:Choice>'
        f'<mc:Fallback><p:transition spd="slow" advTm="{adv_ms}"><p:fade/></p:transition></mc:Fallback>'
        f'</mc:AlternateContent>'
    )


def _insert_transition(slide, el):
    sld = slide._element
    for old in sld.findall(qn("p:transition")) + sld.findall(f"{{{NS_MC}}}AlternateContent"):
        sld.remove(old)
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = sld.find(qn("p:cSld"))
    anchor.addnext(el)


def _shadow(shape):
    spPr = shape._element.spPr
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="50800", dist="25400",
                          dir="5400000", algn="tl", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
    etree.SubElement(clr, qn("a:alpha"), val="15000")


def _add_box(slide, b):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if b["radius"] > 0 else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(kind, Inches(b["x"]), Inches(b["y"]),
                                 Inches(b["w"]), Inches(b["h"]))
    shp.name = "!!" + b["id"]  # '!!' 접두어 = 모핑 강제 매칭
    if b["fill"]:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(b["fill"])
        if b["opacity"] < 1:
            clr = shp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha"), val=str(int(b["opacity"] * 100000)))
    else:
        shp.fill.background()
    shp.line.fill.background()
    if b["radius"] > 0:
        shp.adjustments[0] = max(0.0, min(0.5, b["radius"] / min(b["w"], b["h"])))
    if b["shadow"]:
        _shadow(shp)
    # 기본 도형 스타일(테마 그림자/글꼴색) 제거
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)

    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    pad = Inches(b["pad"])
    tf.margin_left = tf.margin_right = pad
    tf.margin_top = tf.margin_bottom = Inches(min(b["pad"], 0.2))
    tf.vertical_anchor = _ANCHOR[b["anchor"]]
    for i, ln in enumerate(b["lines"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[b["align"]]
        p.space_after = Pt(ln["gap"]) if i < len(b["lines"]) - 1 else Pt(0)
        r = p.add_run()
        r.text = ln["t"]
        f = r.font
        f.size = Pt(ln["size"])
        f.bold = ln["bold"]
        f.name = FONT
        f.color.rgb = RGBColor.from_string(ln["color"])
        rPr = r._r.get_or_add_rPr()
        etree.SubElement(rPr, qn("a:ea"), typeface=FONT)
    return shp


def build_pptx(meta, scenes, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    layout = prs.slide_layouts[6]
    logo = meta.get("logo")
    for i, sc in enumerate(scenes):
        slide = prs.slides.add_slide(layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(TOKENS["bg"])
        for b in sc["boxes"]:
            _add_box(slide, b)
        if logo and Path(logo).exists():  # 로고 미지정 시 우측 상단 비움
            pic = slide.shapes.add_picture(str(logo), Inches(10.13), Inches(0.88),
                                           height=Inches(0.39))
            pic.left = Inches(12.73) - pic.width
            pic.name = "!!_logo"
        kind = "morph" if i > 0 else "fade"
        _insert_transition(slide, _transition_xml(kind, int(sc["morph"] * 1000),
                                                  int(sc["hold"] * 1000)))
        if sc["notes"]:
            slide.notes_slide.notes_text_frame.text = sc["notes"]
    prs.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# ② 애니메이션 HTML (결정론적 seek(t) — 캡처·미리보기 공용)
# ---------------------------------------------------------------------------

def _timeline(meta, scenes):
    t, segs = 0.0, []
    for sc in scenes:
        m0 = t
        t += sc["morph"]
        h0 = t
        t += sc["hold"]
        segs.append(dict(m0=m0, h0=h0, h1=t))
    return segs, t + float(meta.get("outro", 0))


def build_html(meta, scenes, out_path, embed_font=True):
    segs, duration = _timeline(meta, scenes)
    ids = []
    for sc in scenes:
        for b in sc["boxes"]:
            if b["id"] not in ids:
                ids.append(b["id"])
    boxes = {}
    layers = []  # (box id, 텍스트 키) 단위 — 같은 문구면 크기·색 보간, 다르면 교차 페이드
    for eid in ids:
        states, lkeys = [], {}
        for si, sc in enumerate(scenes):
            b = next((x for x in sc["boxes"] if x["id"] == eid), None)
            states.append(None if b is None else {k: b[k] for k in (
                "x", "y", "w", "h", "fill", "radius", "shadow", "opacity")})
            if b is not None and b["lines"]:
                key = "␞".join(l["t"] for l in b["lines"])
                lkeys.setdefault(key, [None] * len(scenes))[si] = dict(
                    pad=b["pad"], align=b["align"], anchor=b["anchor"], lines=b["lines"])
        boxes[eid] = states
        for key, st in lkeys.items():
            layers.append(dict(box=eid, states=st))
    data = dict(W=SLIDE_W, H=SLIDE_H, PX=PX_PER_IN, bg=TOKENS["bg"], segs=segs,
                duration=duration, ids=ids, boxes=boxes, layers=layers)

    font_css = ""
    if embed_font and FONT_FILE.exists():
        b64 = base64.b64encode(FONT_FILE.read_bytes()).decode()
        font_css = ("@font-face{font-family:'Pretendard';font-weight:45 920;"
                    f"src:url(data:font/woff2;base64,{b64}) format('woff2-variations');}}")
    html = HTML_TEMPLATE.replace("/*FONT*/", font_css).replace(
        "/*DATA*/", json.dumps(data, ensure_ascii=False)).replace(
        "__TITLE__", meta.get("project_title") or "Motion Slides")
    Path(out_path).write_text(html, encoding="utf-8")
    return out_path, duration


HTML_TEMPLATE = r"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITLE__</title>
<style>
/*FONT*/
html,body{margin:0;background:#111;height:100%}
body{display:flex;align-items:center;justify-content:center;overflow:hidden}
#stage{position:relative;width:1920px;height:1080px;overflow:hidden;transform-origin:center;
  font-family:'Pretendard','Malgun Gothic','Apple SD Gothic Neo',sans-serif;
  -webkit-font-smoothing:antialiased;font-feature-settings:'tnum'}
.box{position:absolute;box-sizing:border-box}
.tl{position:absolute;inset:0;display:flex;flex-direction:column;box-sizing:border-box}
.tl p{margin:0;line-height:1.2;white-space:pre-wrap;word-break:keep-all}
body.capture{background:transparent;display:block}
body.capture #stage{transform:none!important}
</style></head>
<body><div id="stage"></div>
<script>
const D=/*DATA*/;
const S=document.getElementById('stage'), PX=D.PX, PT=PX/72;
S.style.background='#'+D.bg;
const hex=h=>h?[parseInt(h.slice(0,2),16),parseInt(h.slice(2,4),16),parseInt(h.slice(4,6),16),1]:null;
const lerp=(a,b,p)=>a+(b-a)*p;
const ease=p=>p<.5?4*p*p*p:1-Math.pow(-2*p+2,3)/2;   // 모핑 유사 cubic in-out
const fillOf=(a,b)=>{const A=hex(a&&a.fill),B=hex(b&&b.fill);
  if(!A&&!B)return null; return [A||[...B.slice(0,3),0],B||[...A.slice(0,3),0]];};
const els={}, lels=[];
for(const id of D.ids){const d=document.createElement('div');d.className='box';S.appendChild(d);els[id]=d;}
for(const L of D.layers){
  const d=document.createElement('div');d.className='tl';els[L.box].appendChild(d);
  const ref=L.states.find(s=>s);
  for(const ln of ref.lines){const p=document.createElement('p');p.textContent=ln.t;d.appendChild(p);}
  lels.push(d);
}
// 시점 t → (장면 a, 장면 b, 진행률 p)
function locate(t){
  const g=D.segs;
  for(let i=0;i<g.length;i++){
    if(t<g[i].h0){const p=(t-g[i].m0)/Math.max(1e-6,g[i].h0-g[i].m0);return [i-1,i,ease(Math.min(1,Math.max(0,p)))];}
    if(t<g[i].h1)return [i,i,1];
  }
  return [g.length-1,g.length-1,1];
}
function seek(t){
  const [a,b,p]=locate(t);
  for(const id of D.ids){
    const A=a>=0?D.boxes[id][a]:null, B=D.boxes[id][b], d=els[id];
    const s=B||A; let op;
    if(A&&B)op=lerp(A.opacity,B.opacity,p); else if(B)op=B.opacity*p; else if(A)op=A.opacity*(1-p); else {d.style.display='none';continue;}
    d.style.display='block';
    const g=(k)=>A&&B?lerp(A[k],B[k],p):s[k];
    d.style.left=g('x')*PX+'px'; d.style.top=g('y')*PX+'px';
    d.style.width=g('w')*PX+'px'; d.style.height=g('h')*PX+'px';
    d.style.borderRadius=g('radius')*PX+'px'; d.style.opacity=op;
    const F=fillOf(A&&B?A:s,A&&B?B:s);
    if(F){const c=[0,1,2,3].map(k=>lerp(F[0][k],F[1][k],p));d.style.background=`rgba(${c[0]|0},${c[1]|0},${c[2]|0},${c[3]})`;}
    else d.style.background='transparent';
    const sh=A&&B?lerp(+A.shadow,+B.shadow,p):+s.shadow;
    d.style.boxShadow=sh?`0 ${3*sh}px ${8*sh}px rgba(0,0,0,${.15*sh})`:'none';
  }
  D.layers.forEach((L,i)=>{
    const A=a>=0?L.states[a]:null, B=L.states[b], d=lels[i], s=B||A;
    if(!s){d.style.opacity=0;return;}
    // 문구 교체 시 겹침 최소화: 이전 문구 먼저 사라진 뒤 새 문구 등장
    d.style.opacity=A&&B?1:B?(a<0?p:Math.max(0,(p-.5)/.5)):Math.max(0,1-p/.45);
    d.style.padding=`${Math.min(s.pad,.2)*PX}px ${s.pad*PX}px`;
    d.style.justifyContent={top:'flex-start',middle:'center',bottom:'flex-end'}[s.anchor];
    d.style.textAlign=s.align;
    [...d.children].forEach((el,j)=>{
      const la=A&&A.lines[j], lb=B&&B.lines[j], l=lb||la;
      const sz=la&&lb?lerp(la.size,lb.size,p):l.size;
      const ca=hex(la?la.color:l.color), cb=hex(lb?lb.color:l.color), c=[0,1,2].map(k=>lerp(ca[k],cb[k],(la&&lb)?p:1));
      el.style.fontSize=sz*PT+'px'; el.style.fontWeight=l.bold?700:400;
      el.style.color=`rgb(${c[0]|0},${c[1]|0},${c[2]|0})`;
      el.style.marginBottom=(j<d.children.length-1?l.gap*PT:0)+'px';
    });
  });
}
window.__seek=seek; window.__duration=D.duration;
// 미리보기 재생 (캡처 모드는 ?capture)
if(location.search.includes('capture')){document.body.classList.add('capture');seek(0);}
else{
  const fit=()=>{const k=Math.min(innerWidth/1920,innerHeight/1080);S.style.transform=`scale(${k})`;};
  addEventListener('resize',fit);fit();
  let t0=null;const loop=ts=>{if(t0===null)t0=ts;seek(((ts-t0)/1000)%D.duration);requestAnimationFrame(loop);};
  document.fonts.ready.then(()=>requestAnimationFrame(loop));
}
</script></body></html>
"""


# ---------------------------------------------------------------------------
# ③ MP4 (headless Chromium 프레임 캡처 → FFmpeg)
# ---------------------------------------------------------------------------

def _launch(pw):
    """Chromium 실행 — MOTION_CHROMIUM 환경변수 → Playwright 기본 → 사전 설치본 순"""
    err = None
    for exe in [os.environ.get("MOTION_CHROMIUM"), None, "/opt/pw-browsers/chromium"]:
        if exe is not None and not (exe and Path(exe).exists()):
            continue
        try:
            return pw.chromium.launch(**({"executable_path": exe} if exe else {}))
        except Exception as e:  # 버전 불일치 등 → 다음 후보
            err = e
    raise RuntimeError(f"Chromium 실행 실패 — setup.sh 확인 필요\n{err}")


def render_mp4(html_path, out_path, duration=None, fps=30, crf=18):
    """duration 미지정 시 페이지의 window.__duration 사용 (비트 엔진)"""
    from playwright.sync_api import sync_playwright

    ffmpeg = shutil.which("ffmpeg") or "ffmpeg"
    cmd = [ffmpeg, "-y", "-loglevel", "error", "-f", "image2pipe", "-framerate", str(fps),
           "-i", "-", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-crf", str(crf),
           "-preset", "medium", "-movflags", "+faststart", str(out_path)]
    with sync_playwright() as pw:
        browser = _launch(pw)
        page = browser.new_page(viewport={"width": 1920, "height": 1080})
        page.goto(Path(html_path).resolve().as_uri() + "?capture")
        page.evaluate("document.fonts.ready")
        if duration is None:
            duration = float(page.evaluate("window.__duration"))
        frames = int(round(duration * fps))
        proc = subprocess.Popen(cmd, stdin=subprocess.PIPE)
        try:
            for f in range(frames):
                page.evaluate(f"window.__seek({f / fps})")
                # JPEG(q95) 파이프 — PNG 대비 프레임당 캡처 시간 최대 1/6
                proc.stdin.write(page.screenshot(type="jpeg", quality=95))
        finally:
            proc.stdin.close()
            proc.wait()
        browser.close()
    if proc.returncode:
        raise RuntimeError("ffmpeg 인코딩 실패")
    return out_path, frames
